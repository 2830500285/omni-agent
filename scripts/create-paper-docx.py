from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: create-paper-docx.py <paper.md> <output.docx>", file=sys.stderr)
        return 2

    source = Path(sys.argv[1])
    target = Path(sys.argv[2])
    markdown = source.read_text(encoding="utf-8")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.name = "Times New Roman"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.name = "Times New Roman"
    styles["Heading 2"].font.size = Pt(12)

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            index += 1
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_runs(paragraph, line[2:].strip())
            index += 1
            continue

        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            index += 1
            continue

        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            index += 1
            continue

        if is_table_start(lines, index):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue

        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, line[2:].strip())
            index += 1
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            next_line = lines[index].rstrip()
            if not next_line or next_line.startswith(("# ", "## ", "### ", "- ")) or next_line.strip().startswith("|"):
                break
            paragraph_lines.append(next_line)
            index += 1
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Inches(0.18)
        paragraph.paragraph_format.space_after = Pt(5)
        add_inline_runs(paragraph, " ".join(part.strip() for part in paragraph_lines))

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    print(target)
    return 0


def is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    return lines[index].strip().startswith("|") and re.match(r"^\|?[\s:\-|\|]+\|?$", lines[index + 1].strip()) is not None


def add_markdown_table(document: Document, table_lines: list[str]) -> None:
    rows = [split_table_row(line) for line in table_lines]
    rows = [row for idx, row in enumerate(rows) if idx != 1]
    if not rows:
        return
    table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, cell_text in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
            add_inline_runs(paragraph, cell_text.strip())
    document.add_paragraph()


def split_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def add_inline_runs(paragraph, text: str) -> None:
    cleaned = text.replace("**", "")
    cleaned = re.sub(r"`([^`]+)`", r"\1", cleaned)
    paragraph.add_run(cleaned)


if __name__ == "__main__":
    raise SystemExit(main())
